#!/usr/bin/env python3
"""Render a DiagramSpec (documents/<id>/topology.json) to Mermaid, D2, SVG and HTML (#485 gap 2).

    damira diagram render documents/<id>/topology.json [--pptx] [--redact --out-dir DIR]

Writes next to the spec (or into --out-dir):
  topology.mmd   Mermaid `graph TD`, for pasting into docs and wikis
  topology.d2    D2, for people who keep diagrams as code
  topology.svg   a deterministic tiered layout: tiers top to bottom, a barycenter sort to
                 cut crossings, role shapes and colours, interface and subnet labels.
                 No fonts, images or links are fetched, so it renders offline anywhere.
  index.html     the SVG inline plus node and link tables; self-contained, no scripts
  topology.pptx  only with --pptx: one editable slide of native shapes (python-pptx)

Everything except the pptx is stdlib. python-pptx is found in this interpreter, else via
`uv run --with python-pptx` (no pip install), else skipped with exit code 4: the other
files are still written.

--redact is for anything that leaves the machine (an artifact, a canvas, a ticket):
hostnames, AS labels, management IPs and subnets become stable tokens (HOST-1, IP-2,
NET-1) across every file, using the redact_spec helpers. It never writes the pptx or a
copy of the spec. The same spec always gives byte-identical files.
"""

from __future__ import annotations

import argparse
import html
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import topology  # noqa: E402 — sibling, stdlib only

EXIT_OK = 0
EXIT_ERROR = 1
EXIT_NO_PPTX = 4
_REEXEC_ENV = "DAMIRA_DIAGRAM_CHILD"
_CWD_ENV = "DAMIRA_DIAGRAM_CWD"
_PPTX_PIN = "python-pptx==1.0.2"
LAYOUTS = ("auto", "tiered", "hub-spoke", "leaf-spine")

# Ported from .claude/skills/network-diagrams-pptx so slides and SVGs match.
ROLE_COLORS = {
    "router": "#0096D6", "switch": "#00B48A", "firewall": "#FF4560", "wlc": "#00D4FF",
    "ap": "#00D4FF", "server": "#8B5CF6", "host": "#6B7B93", "cloud": "#FFB800",
}
LINK_STYLES = {  # colour, dashed
    "l2": ("#4A5A73", False), "l3": ("#0096D6", False), "bgp": ("#FFB800", True), "ospf": ("#00D4FF", True),
}
BG, FG, MUTED, LABEL_BG = "#0B1220", "#E2E8F0", "#94A3B8", "#111A2E"
SANS = "system-ui, -apple-system, 'Segoe UI', Roboto, sans-serif"
MONO = "ui-monospace, SFMono-Regular, Menlo, Consolas, monospace"
NODE_W, NODE_H, HGAP, VGAP, MARGIN, TITLE_H = 156, 56, 56, 150, 40, 56


class SpecError(Exception):
    pass


# ---------------------------------------------------------------------------
# Spec
# ---------------------------------------------------------------------------

def normalise(data) -> dict:
    """Validate what the parser or the host model wrote; fill defaults. Raises SpecError."""
    if not isinstance(data, dict):
        raise SpecError("diagram spec must be a JSON object")
    raw_nodes = data.get("nodes")
    if not isinstance(raw_nodes, list) or not raw_nodes:
        raise SpecError("diagram spec needs a non-empty 'nodes' list")
    nodes, seen = [], set()
    for i, n in enumerate(raw_nodes):
        if not isinstance(n, dict) or not str(n.get("id") or "").strip():
            raise SpecError(f"node {i + 1} needs an 'id'")
        nid = str(n["id"]).strip()
        if nid in seen:
            raise SpecError(f"node id {nid!r} appears twice")
        seen.add(nid)
        role = str(n.get("role") or "").lower()
        node = {
            "id": nid, "label": str(n.get("label") or nid),
            "role": role if role in topology.ROLES else "router",
            "platform": str(n.get("platform") or ""), "mgmt_ip": str(n.get("mgmt_ip") or ""),
            "zone": str(n.get("zone") or ""), "external": bool(n.get("external")),
        }
        tier = n.get("tier")
        node["tier"] = tier if isinstance(tier, int) and tier >= 0 else topology.infer_tier(node)
        nodes.append(node)
    links = []
    for i, lk in enumerate(data.get("links") or []):
        if not isinstance(lk, dict):
            continue
        a, b = str(lk.get("a") or ""), str(lk.get("b") or "")
        for end in (a, b):
            if end not in seen:
                raise SpecError(f"link {i + 1} names {end!r}, which is not a node")
        kind = str(lk.get("kind") or "l2").lower()
        links.append({
            "a": a, "a_if": str(lk.get("a_if") or ""), "b": b, "b_if": str(lk.get("b_if") or ""),
            "subnet": str(lk.get("subnet") or ""), "kind": kind if kind in topology.KINDS else "l2",
            "source": str(lk.get("source") or "manual"),
            "protocols": sorted({str(p) for p in (lk.get("protocols") or [])}),
        })
    zones = {}
    for z in data.get("zones") or []:
        if isinstance(z, dict) and z.get("id"):
            zones[str(z["id"])] = str(z.get("label") or z["id"])
        elif isinstance(z, str) and z:
            zones[z] = z
    for n in nodes:
        if n["zone"]:
            zones.setdefault(n["zone"], n["zone"])
    layout = str(data.get("layout") or "auto")
    return {
        "title": str(data.get("title") or "Network topology"),
        "layout": layout if layout in LAYOUTS else "auto",
        "nodes": sorted(nodes, key=lambda n: n["id"]),
        "links": sorted(links, key=lambda lk: (lk["a"], lk["b"], lk["a_if"], lk["b_if"], lk["kind"])),
        "zones": [{"id": k, "label": zones[k]} for k in sorted(zones)],
    }


# ---------------------------------------------------------------------------
# Layout
# ---------------------------------------------------------------------------

def _adjacency(spec: dict) -> "dict[str, set]":
    adj = {n["id"]: set() for n in spec["nodes"]}
    for lk in spec["links"]:
        if lk["a"] != lk["b"]:
            adj[lk["a"]].add(lk["b"])
            adj[lk["b"]].add(lk["a"])
    return adj


def resolve_layout(spec: dict) -> str:
    if spec["layout"] != "auto":
        return spec["layout"]
    names = [n["id"].lower() for n in spec["nodes"]]
    if any("spine" in x for x in names) and any("leaf" in x for x in names):
        return "leaf-spine"
    adj = _adjacency(spec)
    inner = [n["id"] for n in spec["nodes"] if not n["external"]]
    for hub in inner:
        spokes = [x for x in inner if x != hub]
        if len(spokes) >= 3 and all(hub in adj[s] and not (adj[s] & set(spokes)) for s in spokes):
            return "hub-spoke"
    return "tiered"


def tiers(spec: dict) -> "list[list[str]]":
    """Rows of node ids, top to bottom, ordered to reduce crossings."""
    layout = resolve_layout(spec)
    tier = {}
    for n in spec["nodes"]:
        t = n["tier"]
        name = n["id"].lower()
        if layout == "leaf-spine" and not n["external"]:
            t = 100 if "spine" in name else 101 if "leaf" in name else 102 + n["tier"]
        tier[n["id"]] = t
    if layout == "hub-spoke":
        adj = _adjacency(spec)
        inner = [n["id"] for n in spec["nodes"] if not n["external"]]
        hub = max(inner, key=lambda x: (len(adj[x]), x[::-1]))
        for x in inner:
            tier[x] = 100 if x == hub else 101
    levels = sorted(set(tier.values()))
    rows = [sorted(nid for nid, t in tier.items() if t == lv) for lv in levels]

    adj = _adjacency(spec)

    def sweep(order):
        for r in order:
            ref = {nid: (i + 0.5) / len(row) for row in rows for i, nid in enumerate(row)}
            src = rows[r - 1] if order[0] == 1 else rows[r + 1]
            src_set = set(src)

            def key(nid):
                near = [ref[x] for x in adj[nid] if x in src_set]
                return (sum(near) / len(near) if near else ref[nid], nid)
            rows[r] = sorted(rows[r], key=key)

    for _ in range(3):
        sweep(list(range(1, len(rows))))
        sweep(list(range(len(rows) - 2, -1, -1)))
    return rows


def positions(rows: "list[list[str]]") -> "tuple[dict, int, int]":
    widest = max(len(r) for r in rows)
    width = max(520, widest * (NODE_W + HGAP) - HGAP + 2 * MARGIN)
    pos = {}
    for t, row in enumerate(rows):
        row_w = len(row) * (NODE_W + HGAP) - HGAP
        x0 = (width - row_w) / 2
        for i, nid in enumerate(row):
            pos[nid] = (x0 + i * (NODE_W + HGAP) + NODE_W / 2, MARGIN + TITLE_H + t * (NODE_H + VGAP) + NODE_H / 2)
    height = MARGIN + TITLE_H + len(rows) * (NODE_H + VGAP) - VGAP + MARGIN + 44  # + legend
    return pos, int(width), int(height)


def _bez(c, t):
    x1, y1, cx, cy, x2, y2 = c
    u = 1 - t
    return (u * u * x1 + 2 * u * t * cx + t * t * x2, u * u * y1 + 2 * u * t * cy + t * t * y2)


def mid_label(lk: dict) -> str:
    if lk["kind"] in ("bgp", "ospf") and not lk["subnet"]:
        return lk["kind"].upper()
    return " ".join(x for x in (lk["subnet"], "/".join(p for p in lk["protocols"] if p != lk["kind"])) if x)


def _tag_w(text: str, size: int = 10) -> float:
    return len(text) * size * 0.62 + 8


def _anchors(spec: dict, pos: dict) -> dict:
    """Each link end leaves from its own point on the node's edge, spread along the side
    that faces the other end and ordered by where that end is, so links don't cross at the
    node and their port labels don't stack. {(link index, 0|1): (x, y, slot)}."""
    sides: "dict[tuple, list]" = {}
    for i, lk in enumerate(spec["links"]):
        for which, me, other in ((0, lk["a"], lk["b"]), (1, lk["b"], lk["a"])):
            (mx, my), (ox, oy) = pos[me], pos[other]
            side = "top" if oy < my - 1 else "bottom" if oy > my + 1 else ("right" if ox > mx else "left")
            sides.setdefault((me, side), []).append((ox, oy, i, which))
    out = {}
    for (nid, side), items in sides.items():
        items.sort()
        x, y = pos[nid]
        k = len(items)
        for j, (_, _, i, which) in enumerate(items):
            if side in ("top", "bottom"):
                out[(i, which)] = (x - NODE_W / 2 + (j + 1) * NODE_W / (k + 1),
                                   y + (NODE_H / 2 if side == "bottom" else -NODE_H / 2), j)
            else:
                out[(i, which)] = (x + (NODE_W / 2 if side == "right" else -NODE_W / 2),
                                   y - NODE_H / 2 + (j + 1) * NODE_H / (k + 1), j)
    return out


def _hits(box, x, y, pad=6.0) -> bool:
    bx, by = box
    return abs(x - bx) < NODE_W / 2 + pad and abs(y - by) < NODE_H / 2 + pad


def routes(spec: dict, pos: dict) -> "list[dict]":
    """A quadratic curve per link: straight when that is clear, otherwise bent (the least it
    takes) so the line misses every other node and its subnet label sits on no node.
    Returns [{link, c: (x1, y1, cx, cy, x2, y2), ta, tb}], ta/tb being where port labels go."""
    anchors = _anchors(spec, pos)
    out = []
    for i, lk in enumerate(spec["links"]):
        x1, y1, ja = anchors[(i, 0)]
        x2, y2, jb = anchors[(i, 1)]
        dx, dy = x2 - x1, y2 - y1
        length = max((dx * dx + dy * dy) ** 0.5, 1.0)
        px, py = -dy / length, dx / length
        up = -1 if py > 0 else 1  # the offset sign that bows a same-tier link upwards
        steps = [0] + [s * m for m in (50, 100, 150, 200, 260) for s in ((up, -up) if abs(dy) < 1 else (1, -1))]
        others = [pos[n] for n in pos if n not in (lk["a"], lk["b"])]
        label = mid_label(lk)
        chosen = (x1, y1, (x1 + x2) / 2, (y1 + y2) / 2, x2, y2)
        for off in steps:
            c = (x1, y1, (x1 + x2) / 2 + px * off, (y1 + y2) / 2 + py * off, x2, y2)
            if any(_hits(b, *_bez(c, t / 25)) for b in others for t in range(1, 25)):
                continue
            if label:
                mx, my = _bez(c, 0.5)
                half_w = _tag_w(label) / 2
                if any(abs(mx - bx) < NODE_W / 2 + half_w + 4 and abs(my - by) < NODE_H / 2 + 12
                       for bx, by in pos.values()):
                    continue
            chosen = c
            break
        arc = length + abs(chosen[3] - (y1 + y2) / 2) + abs(chosen[2] - (x1 + x2) / 2)
        ta = min(0.45, (15 + (ja % 2) * 18) / arc)
        tb = 1 - min(0.45, (15 + (jb % 2) * 18) / arc)
        out.append({"link": lk, "c": chosen, "ta": ta, "tb": tb})
    return out


# ---------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------

def _num(v: float) -> str:
    return f"{v:.1f}".rstrip("0").rstrip(".")


def _clip(text: str, n: int) -> str:
    return text if len(text) <= n else text[: n - 1] + "…"


def _shape(role: str, x: float, y: float, fill: str, dashed: bool) -> str:
    w, h = NODE_W, NODE_H
    left, top = x - w / 2, y - h / 2
    if dashed:  # external: a see-through fill over a solid backing, so lines don't show through
        return _shape(role, x, y, BG, False).replace(f' stroke="{FG}" stroke-width="1.5"', "") + _shape(
            role, x, y, fill, False).replace('stroke-width="1.5"', 'stroke-width="1.5" stroke-dasharray="5 4" '
                                             'fill-opacity="0.45"')
    stroke = f' stroke="{FG}" stroke-width="1.5"'
    if role == "cloud":
        return f'<ellipse cx="{_num(x)}" cy="{_num(y)}" rx="{_num(w / 2)}" ry="{_num(h / 2)}" fill="{fill}"{stroke}/>'
    if role == "firewall":
        k = 14
        pts = [(left + k, top), (left + w - k, top), (left + w, y), (left + w - k, top + h), (left + k, top + h), (left, y)]
        return f'<polygon points="{" ".join(f"{_num(a)},{_num(b)}" for a, b in pts)}" fill="{fill}"{stroke}/>'
    rx = {"router": 18, "ap": h / 2, "wlc": 10}.get(role, 3)
    return (f'<rect x="{_num(left)}" y="{_num(top)}" width="{w}" height="{h}" rx="{_num(rx)}" '
            f'fill="{fill}"{stroke}/>')


def _tag(x: float, y: float, text: str, color: str, font: str = MONO, size: int = 10) -> str:
    w = len(text) * size * 0.62 + 8
    return (f'<rect x="{_num(x - w / 2)}" y="{_num(y - size / 2 - 3)}" width="{_num(w)}" height="{size + 6}" '
            f'rx="3" fill="{LABEL_BG}" fill-opacity="0.92"/>'
            f'<text x="{_num(x)}" y="{_num(y + size / 2 - 1)}" font-family="{font}" font-size="{size}" '
            f'fill="{color}" text-anchor="middle">{html.escape(text)}</text>')


def render_svg(spec: dict) -> str:
    rows = tiers(spec)
    pos, width, height = positions(rows)
    nodes = {n["id"]: n for n in spec["nodes"]}
    e = html.escape
    out = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" '
           f'viewBox="0 0 {width} {height}" role="img" aria-label="{e(spec["title"])}">',
           f'<rect width="{width}" height="{height}" fill="{BG}"/>',
           f'<text x="{MARGIN}" y="{MARGIN + 8}" font-family="{SANS}" font-size="18" font-weight="600" '
           f'fill="{FG}">{e(spec["title"])}</text>',
           f'<text x="{MARGIN}" y="{MARGIN + 28}" font-family="{SANS}" font-size="11" fill="{MUTED}">'
           f'{len(spec["nodes"])} devices, {len(spec["links"])} links · {resolve_layout(spec)} layout</text>']
    labels = []
    for r in routes(spec, pos):
        lk, c = r["link"], r["c"]
        color, dashed = LINK_STYLES[lk["kind"]]
        out.append(f'<path d="M{_num(c[0])} {_num(c[1])} Q{_num(c[2])} {_num(c[3])} {_num(c[4])} {_num(c[5])}" '
                   f'fill="none" stroke="{color}" stroke-width="{2.5 if lk["kind"] == "l3" else 2}"'
                   + (' stroke-dasharray="7 5"' if dashed else "") + "/>")
        for t, name in ((r["ta"], lk["a_if"]), (r["tb"], lk["b_if"])):
            if name:
                labels.append(_tag(*_bez(c, t), name, FG))
        mid = mid_label(lk)
        if mid:
            labels.append(_tag(*_bez(c, 0.5), mid, color if lk["kind"] != "l2" else MUTED))
    for nid, (x, y) in sorted(pos.items()):
        n = nodes[nid]
        out.append(_shape(n["role"], x, y, ROLE_COLORS[n["role"]], n["external"]))
        sub = n["platform"] or n["role"]
        out.append(f'<text x="{_num(x)}" y="{_num(y - 3)}" font-family="{SANS}" font-size="13" font-weight="600" '
                   f'fill="#FFFFFF" text-anchor="middle">{e(_clip(n["label"], 20))}</text>'
                   f'<text x="{_num(x)}" y="{_num(y + 13)}" font-family="{SANS}" font-size="10" '
                   f'fill="#FFFFFF" fill-opacity="0.8" text-anchor="middle">{e(_clip(sub, 24))}</text>')
    out += labels
    kinds = [k for k in LINK_STYLES if any(lk["kind"] == k for lk in spec["links"])]
    names = {"l2": "L2 (CDP/LLDP)", "l3": "L3 link", "bgp": "BGP session, no link", "ospf": "OSPF, no link"}
    lx, ly = MARGIN, height - MARGIN + 6
    for k in kinds:
        color, dashed = LINK_STYLES[k]
        out.append(f'<line x1="{lx}" y1="{ly}" x2="{lx + 28}" y2="{ly}" stroke="{color}" stroke-width="2.5"'
                   + (' stroke-dasharray="7 5"' if dashed else "") + "/>"
                   f'<text x="{lx + 36}" y="{ly + 4}" font-family="{SANS}" font-size="11" fill="{MUTED}">'
                   f'{names[k]}</text>')
        lx += 36 + len(names[k]) * 7 + 24
    out.append("</svg>")
    return "\n".join(out) + "\n"


def mermaid_id(nid: str) -> str:
    mid = re.sub(r"\W", "_", nid)
    return f"n_{mid}" if not mid or mid[0].isdigit() or mid.lower() in ("end", "graph", "subgraph") else mid


_MERMAID_SHAPES = {"router": ('("', '")'), "switch": ('["', '"]'), "firewall": ('{{"', '"}}'),
                   "cloud": ('(("', '"))'), "ap": ('(["', '"])'), "wlc": ('(["', '"])'),
                   "server": ('[["', '"]]'), "host": ('[["', '"]]')}


def _mq(text: str) -> str:
    return text.replace('"', "#quot;").replace("<", "#lt;").replace(">", "#gt;")


def render_mermaid(spec: dict) -> str:
    out = ["%% Generated by damira diagram render from topology.json. Edit the spec, not this file.",
           "graph TD"]

    def node_line(n, indent):
        o, c = _MERMAID_SHAPES[n["role"]]
        label = _mq(n["label"]) + (f"<br/>{_mq(n['platform'])}" if n["platform"] else "")
        return f"{indent}{mermaid_id(n['id'])}{o}{label}{c}"

    for z in spec["zones"]:
        out.append(f'  subgraph {mermaid_id("zone_" + z["id"])}["{_mq(z["label"])}"]')
        out += [node_line(n, "    ") for n in spec["nodes"] if n["zone"] == z["id"]]
        out.append("  end")
    out += [node_line(n, "  ") for n in spec["nodes"] if not n["zone"]]
    for lk in spec["links"]:
        ports = " — ".join(p for p in (lk["a_if"], lk["b_if"]) if p)
        extra = " ".join(x for x in (lk["subnet"], "/".join(lk["protocols"])) if x)
        label = "<br/>".join(x for x in (_mq(ports), _mq(extra)) if x) or lk["kind"]
        arrow = "-.-" if LINK_STYLES[lk["kind"]][1] else "---"
        out.append(f'  {mermaid_id(lk["a"])} {arrow}|"{label}"| {mermaid_id(lk["b"])}')
    for role in topology.ROLES:
        members = [mermaid_id(n["id"]) for n in spec["nodes"] if n["role"] == role]
        if members:
            out.append(f"  classDef {role} fill:{ROLE_COLORS[role]},color:#fff,stroke:#fff")
            out.append(f"  class {','.join(members)} {role}")
    ext = [mermaid_id(n["id"]) for n in spec["nodes"] if n["external"]]
    if ext:
        out.append("  classDef external stroke-dasharray: 5 4")
        out.append(f"  class {','.join(ext)} external")
    return "\n".join(out) + "\n"


_D2_SHAPES = {"router": "rectangle", "switch": "rectangle", "firewall": "hexagon", "cloud": "cloud",
              "ap": "oval", "wlc": "rectangle", "server": "rectangle", "host": "rectangle"}


def _d2esc(text: str) -> str:
    return text.replace("\\", "\\\\").replace('"', '\\"')


def _dq(text: str) -> str:
    return f'"{_d2esc(text)}"'


def render_d2(spec: dict) -> str:
    out = ["# Generated by damira diagram render from topology.json. Edit the spec, not this file.",
           "direction: down", f"title: {_dq(spec['title'])} {{shape: text; near: top-center}}"]
    path = {}
    zone_label = {z["id"]: z["label"] for z in spec["zones"]}

    def node_def(n):
        label = _d2esc(n["label"]) + (r"\n" + _d2esc(n["platform"]) if n["platform"] else "")
        radius = "; style.border-radius: 12" if n["role"] == "router" else ""
        dash = "; style.stroke-dash: 4" if n["external"] else ""
        return (f"{_dq(n['id'])}: \"{label}\" "
                f"{{shape: {_D2_SHAPES[n['role']]}; style.fill: {_dq(ROLE_COLORS[n['role']])}; "
                f"style.font-color: \"#ffffff\"{radius}{dash}}}")

    for z in spec["zones"]:
        out.append(f"{_dq('zone_' + z['id'])}: {_dq(zone_label[z['id']])} {{")
        for n in spec["nodes"]:
            if n["zone"] == z["id"]:
                out.append("  " + node_def(n))
                path[n["id"]] = f"{_dq('zone_' + z['id'])}.{_dq(n['id'])}"
        out.append("}")
    for n in spec["nodes"]:
        if not n["zone"]:
            out.append(node_def(n))
            path[n["id"]] = _dq(n["id"])
    for lk in spec["links"]:
        ports = " - ".join(p for p in (lk["a_if"], lk["b_if"]) if p)
        label = " ".join(x for x in (ports, lk["subnet"], "/".join(lk["protocols"])) if x) or lk["kind"]
        color, dashed = LINK_STYLES[lk["kind"]]
        style = f"style.stroke: {_dq(color)}" + ("; style.stroke-dash: 4" if dashed else "")
        out.append(f"{path[lk['a']]} -- {path[lk['b']]}: {_dq(label)} {{{style}}}")
    return "\n".join(out) + "\n"


def render_html(spec: dict, svg: str) -> str:
    e = html.escape
    rows = "".join(
        f"<tr><td>{e(n['label'])}</td><td>{e(n['role'])}</td><td>{e(n['platform'])}</td>"
        f"<td>{e(n['mgmt_ip'])}</td><td>{e(n['zone'])}</td><td>{'yes' if n['external'] else ''}</td></tr>"
        for n in spec["nodes"])
    label = {n["id"]: n["label"] for n in spec["nodes"]}
    links = "".join(
        f"<tr><td>{e(label[lk['a']])}</td><td>{e(lk['a_if'])}</td><td>{e(label[lk['b']])}</td><td>{e(lk['b_if'])}</td>"
        f"<td>{e(lk['subnet'])}</td><td>{e(lk['kind'])}</td><td>{e(lk['source'])}</td>"
        f"<td>{e(', '.join(lk['protocols']))}</td></tr>"
        for lk in spec["links"])
    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{e(spec['title'])}</title>
<style>
:root {{ color-scheme: dark; }}
body {{ margin: 0; padding: 24px 16px; background: {BG}; color: {FG}; font: 14px/1.5 {SANS}; }}
main {{ max-width: 1200px; margin: 0 auto; }}
h1 {{ font-size: 20px; margin: 0 0 4px; }} h2 {{ font-size: 15px; margin: 28px 0 8px; color: {MUTED}; }}
.note {{ color: {MUTED}; font-size: 12px; margin: 0 0 16px; }}
.diagram {{ overflow-x: auto; border: 1px solid #1E293B; border-radius: 8px; }}
.diagram svg {{ display: block; max-width: 100%; height: auto; }}
.tbl {{ overflow-x: auto; }}
table {{ border-collapse: collapse; width: 100%; font-size: 13px; }}
th, td {{ text-align: left; padding: 6px 10px; border-bottom: 1px solid #1E293B; white-space: nowrap; }}
th {{ color: {MUTED}; font-weight: 600; }} td:nth-child(2), td:nth-child(4) {{ font-family: {MONO}; }}
</style></head>
<body><main>
<h1>{e(spec['title'])}</h1>
<p class="note">Built from saved configs and CDP/LLDP output. Links marked subnet or bgp are inferred; check them against the network before relying on this diagram.</p>
<div class="diagram">
{svg}</div>
<h2>Devices ({len(spec['nodes'])})</h2>
<div class="tbl"><table><thead><tr><th>Device</th><th>Role</th><th>Platform</th><th>Mgmt IP</th><th>Zone</th><th>External</th></tr></thead>
<tbody>{rows}</tbody></table></div>
<h2>Links ({len(spec['links'])})</h2>
<div class="tbl"><table><thead><tr><th>A</th><th>A port</th><th>B</th><th>B port</th><th>Subnet</th><th>Kind</th><th>Source</th><th>Protocols</th></tr></thead>
<tbody>{links}</tbody></table></div>
</main></body></html>
"""


# ---------------------------------------------------------------------------
# pptx (optional, python-pptx)
# ---------------------------------------------------------------------------

def render_pptx(spec: dict, target: Path) -> None:
    """One slide of native, editable shapes, following the network-diagrams-pptx skill."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    def rgb(hexcolor: str):
        return RGBColor.from_string(hexcolor.lstrip("#").upper())

    shapes = {"router": MSO_SHAPE.ROUNDED_RECTANGLE, "switch": MSO_SHAPE.RECTANGLE, "firewall": MSO_SHAPE.HEXAGON,
              "cloud": MSO_SHAPE.CLOUD, "ap": MSO_SHAPE.ROUNDED_RECTANGLE, "wlc": MSO_SHAPE.ROUNDED_RECTANGLE,
              "server": MSO_SHAPE.RECTANGLE, "host": MSO_SHAPE.OVAL}
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BG)

    def text(x, y, w, h, value, size, color, bold=False):
        tb = slide.shapes.add_textbox(x, y, w, h)
        p = tb.text_frame.paragraphs[0]
        p.text, p.alignment = value, PP_ALIGN.CENTER
        p.font.size, p.font.bold, p.font.name = Pt(size), bold, "Calibri"
        p.font.color.rgb = rgb(color)
        return tb

    title = text(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.6), spec["title"], 24, FG, True)
    title.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT

    rows = tiers(spec)
    pos, width, height = positions(rows)
    area_x, area_y, area_w, area_h = Inches(0.5), Inches(1.2), Inches(12.3), Inches(5.9)
    scale = min(area_w / width, area_h / height)
    off_x = area_x + (area_w - width * scale) / 2

    def at(x, y):
        return Emu(int(off_x + x * scale)), Emu(int(area_y + y * scale))

    w, h = Emu(int(NODE_W * scale)), Emu(int(NODE_H * scale))
    nodes = {n["id"]: n for n in spec["nodes"]}
    for r in routes(spec, pos):  # links first so devices sit on top
        lk, c = r["link"], r["c"]
        color, dashed = LINK_STYLES[lk["kind"]]
        if c[2:4] == ((c[0] + c[4]) / 2, (c[1] + c[5]) / 2):
            line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *at(c[0], c[1]), *at(c[4], c[5]))
        else:  # the same bend as the SVG, as an editable polyline
            pts = [at(*_bez(c, t / 16)) for t in range(17)]
            fb = slide.shapes.build_freeform(*pts[0])
            fb.add_line_segments(pts[1:], close=False)
            line = fb.convert_to_shape()
            line.fill.background()
        line.line.color.rgb, line.line.width = rgb(color), Pt(2)
        if dashed:
            line.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        mid = mid_label(lk)
        if mid:
            mx, my = at(*_bez(c, 0.5))
            text(mx - Inches(0.8), my - Inches(0.14), Inches(1.6), Inches(0.28), mid, 8, color)
        for t, name in ((r["ta"], lk["a_if"]), (r["tb"], lk["b_if"])):
            if name:
                px, py = at(*_bez(c, t))
                text(px - Inches(0.5), py - Inches(0.12), Inches(1.0), Inches(0.24), name, 7, MUTED)
    for nid, (x, y) in sorted(pos.items()):
        n = nodes[nid]
        left, top = at(x - NODE_W / 2, y - NODE_H / 2)
        shp = slide.shapes.add_shape(shapes[n["role"]], left, top, w, h)
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(ROLE_COLORS[n["role"]])
        shp.line.color.rgb, shp.line.width = rgb("#FFFFFF"), Pt(1.5)
        if n["external"]:
            shp.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        tf = shp.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text, p.alignment = n["label"], PP_ALIGN.CENTER
        p.font.size, p.font.bold, p.font.name = Pt(11), True, "Calibri"
        p.font.color.rgb = rgb("#FFFFFF")
        if n["platform"]:
            q = tf.add_paragraph()
            q.text, q.alignment = _clip(n["platform"], 24), PP_ALIGN.CENTER
            q.font.size, q.font.name = Pt(8), "Calibri"
            q.font.color.rgb = rgb("#E2E8F0")
    target.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(target))


def _pptx_available() -> bool:
    return importlib.util.find_spec("pptx") is not None


def _reexec_pptx(spec_path: Path, out_dir: "Path | None") -> "int | None":
    uv = shutil.which("uv")
    if not uv or os.environ.get(_REEXEC_ENV):
        return None
    print(f"python-pptx is not installed; rendering the slide with uv (fetches {_PPTX_PIN} once).",
          file=sys.stderr)
    # Same isolation as render_workbook: no project config picks the interpreter or index.
    # This interpreter first; if uv rejects it (some Homebrew builds fail uv's probe), let
    # uv choose one. The slide is idempotent, so a second attempt is harmless.
    script = ["python", str(Path(__file__).resolve()), str(spec_path.resolve()), "--pptx-only"]
    if out_dir:
        script += ["--out-dir", str(out_dir.resolve())]
    env = {**os.environ, _REEXEC_ENV: "1", _CWD_ENV: os.getcwd()}
    code = None
    for pin in (["--python", sys.executable], []):
        cmd = [uv, "run", "--quiet", "--no-project", "--no-config", *pin, "--with", _PPTX_PIN, *script]
        try:
            with tempfile.TemporaryDirectory(prefix="damira-uv-") as neutral:
                code = subprocess.run(cmd, cwd=neutral, env=env, timeout=600).returncode
        except (OSError, subprocess.SubprocessError):
            continue
        if code == 0:
            return 0
    return code


# ---------------------------------------------------------------------------
# Redaction
# ---------------------------------------------------------------------------

def redact_diagram(spec: dict) -> "tuple[dict, dict]":
    """Tokenise hosts, AS labels, IPs, subnets and zones; the same value gets the same token
    in every output. Returns (redacted copy, {kind: count})."""
    import redact_spec

    t = redact_spec._Tokens()
    exact = {}
    for n in spec["nodes"]:
        exact[n["id"]] = t.token(n["id"], "HOST")
    for n in spec["nodes"]:
        if n["label"] not in exact:
            exact[n["label"]] = n["label"] if n["label"] == n["id"] else t.token(n["label"], "HOST")
    for z in spec["zones"]:
        exact.setdefault(z["id"], t.token(z["id"], "ZONE"))
        exact.setdefault(z["label"], t.token(z["label"], "ZONE"))

    def tok(value: str, kind: str) -> str:
        return t.token(value, kind) if value else ""

    def scrub(value: str) -> str:
        return redact_spec._scrub(value, t) if value else value

    nodes = [{**n, "id": exact[n["id"]], "label": exact.get(n["label"], n["label"]),
              "mgmt_ip": tok(n["mgmt_ip"], "IP"), "zone": exact.get(n["zone"], ""),
              "platform": scrub(n["platform"])} for n in spec["nodes"]]
    links = [{**lk, "a": exact[lk["a"]], "b": exact[lk["b"]], "subnet": tok(lk["subnet"], "NET"),
              "a_if": scrub(lk["a_if"]), "b_if": scrub(lk["b_if"])} for lk in spec["links"]]
    out = {**spec, "title": scrub(spec["title"]),
           "zones": [{"id": exact[z["id"]], "label": exact[z["label"]]} for z in spec["zones"]]}
    out["nodes"] = sorted(nodes, key=lambda n: n["id"])
    out["links"] = sorted(links, key=lambda lk: (lk["a"], lk["b"], lk["a_if"], lk["b_if"], lk["kind"]))
    return out, dict(t.counts)


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------

def _stem(spec_path: Path) -> str:
    stem = spec_path.name
    for suffix in (".damira.json", ".json"):
        if stem.endswith(suffix):
            return stem[: -len(suffix)] or "topology"
    return stem or "topology"


def _write(target: Path, data: str) -> None:
    fd, tmp = tempfile.mkstemp(prefix=".damira-", suffix=target.suffix, dir=str(target.parent))
    try:
        with os.fdopen(fd, "w", encoding="utf-8", newline="\n") as fh:
            fh.write(data)
        umask = os.umask(0)
        os.umask(umask)
        os.chmod(tmp, 0o666 & ~umask)
        os.replace(tmp, str(target))
    except BaseException:
        try:
            os.unlink(tmp)
        except OSError:
            pass
        raise


def load(spec_path) -> dict:
    try:
        return normalise(json.loads(Path(spec_path).read_text(encoding="utf-8")))
    except (OSError, json.JSONDecodeError) as exc:
        raise SpecError(f"cannot read {spec_path}: {exc}")


def render(spec_path, out_dir=None, redact=False, pptx=False, host="cli") -> dict:
    """Write .mmd, .d2, .svg and index.html; the pptx too if asked and python-pptx imports."""
    started = time.monotonic()
    spec_path = Path(spec_path)
    spec = load(spec_path)
    redactions = {}
    if redact:
        spec, redactions = redact_diagram(spec)
        pptx = False
    folder = Path(out_dir) if out_dir else spec_path.parent
    folder.mkdir(parents=True, exist_ok=True)
    stem = _stem(spec_path)
    svg = render_svg(spec)
    paths = {"svg": folder / f"{stem}.svg", "mmd": folder / f"{stem}.mmd", "d2": folder / f"{stem}.d2",
             "html": folder / "index.html"}
    _write(paths["svg"], svg)
    _write(paths["mmd"], render_mermaid(spec))
    _write(paths["d2"], render_d2(spec))
    _write(paths["html"], render_html(spec, svg))
    paths["pptx_target"] = folder / f"{stem}.pptx"
    paths["pptx"] = None
    if pptx and _pptx_available():
        render_pptx(spec, paths["pptx_target"])
        paths["pptx"] = paths["pptx_target"]
    status = "redacted" if redact else ("partial" if pptx and not paths["pptx"] else "ok")
    _log(host, status, spec, started)
    return {**paths, "spec": spec, "redactions": redactions}


def _log(host: str, status: str, spec: dict, started: float) -> None:
    try:
        import events
        events.emit({"event": "deliverable_generated", "deliverable_type": "diagram", "host": host,
                     "status": status, "duration_ms": int((time.monotonic() - started) * 1000),
                     "counts": {"nodes": len(spec["nodes"]), "links": len(spec["links"])}})
    except Exception:  # noqa: BLE001 — logging never breaks rendering
        pass


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(prog="damira diagram render",
                                description="Render topology.json to Mermaid, D2, SVG and index.html")
    p.add_argument("spec", help="DiagramSpec JSON, e.g. documents/<id>/topology.json")
    p.add_argument("--out-dir", default="", help="default: the spec's folder")
    p.add_argument("--pptx", action="store_true", help="also write topology.pptx (python-pptx, via uv if needed)")
    p.add_argument("--redact", action="store_true",
                   help="tokenise hosts, IPs and subnets for sharing; use with --out-dir. Never writes a pptx")
    p.add_argument("--pptx-only", action="store_true", help=argparse.SUPPRESS)
    p.add_argument("--host", default=os.environ.get("DAMIRA_HOST", "cli"),
                   choices=["cli", "claude-code", "cursor"], help=argparse.SUPPRESS)
    return p


def main(argv: "list[str] | None" = None) -> int:
    argv = list(sys.argv[1:] if argv is None else argv)
    if os.environ.get(_REEXEC_ENV) and os.environ.get(_CWD_ENV):
        try:
            os.chdir(os.environ[_CWD_ENV])
        except OSError:
            pass
    a = build_parser().parse_args(argv)
    out_dir = Path(a.out_dir) if a.out_dir else None
    try:
        if a.pptx_only:
            target = (out_dir or Path(a.spec).parent) / f"{_stem(Path(a.spec))}.pptx"
            render_pptx(load(a.spec), target)
            print(f"Slide:    {target}")
            return EXIT_OK
        result = render(a.spec, out_dir, redact=a.redact, pptx=a.pptx, host=a.host)
    except SpecError as exc:
        print(f"damira diagram: {exc}", file=sys.stderr)
        return EXIT_ERROR
    except OSError as exc:
        print(f"damira diagram: cannot write the output: {exc.strerror or exc}", file=sys.stderr)
        return EXIT_ERROR

    print(f"SVG:      {result['svg']}")
    print(f"Preview:  {result['html']}")
    print(f"Mermaid:  {result['mmd']}")
    print(f"D2:       {result['d2']}")
    if a.redact:
        done = ", ".join(f"{n} {k}" for k, n in sorted(result["redactions"].items())) or "nothing matched"
        print(f"Redacted: {done}. Check the preview for anything the patterns missed before sharing it.")
        return EXIT_OK
    if a.pptx and not result["pptx"]:
        code = _reexec_pptx(Path(a.spec), out_dir)
        if code == 0:
            return EXIT_OK
        print("No .pptx written: python-pptx isn't installed and uv couldn't provide it. Install uv "
              "(https://docs.astral.sh/uv/) and run this again; the SVG, HTML, Mermaid and D2 are ready.",
              file=sys.stderr)
        return EXIT_NO_PPTX
    if result["pptx"]:
        print(f"Slide:    {result['pptx']}")
    return EXIT_OK


if __name__ == "__main__":
    sys.exit(main())
