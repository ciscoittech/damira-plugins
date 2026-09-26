#!/usr/bin/env python3
# /// script
# requires-python = ">=3.9"
# dependencies = ["openpyxl==3.1.5"]
# ///
"""Render a WorkbookSpec (documents/<id>/workbook.json) to workbook.xlsx + index.html (#484).

The plugin's renderer is the source of truth for the workbook. Artifacts, canvases and
the xlsx skill are presentation layers on top of what this writes.

    damira workbook documents/<id>/workbook.json          # via the damira CLI
    uv run render_workbook.py documents/<id>/workbook.json  # directly (PEP 723)

Pipeline: load spec → grounding check against documents/<id>/evidence/ (tags command
cells UNVERIFIED, never rewrites them) → index.html (stdlib, always) → workbook.xlsx.

The xlsx needs openpyxl, and the plugin stays stdlib-only, so this file imports openpyxl
lazily and finds it in this order:
  1. already importable in this interpreter → render in-process
  2. `uv` on PATH → re-run this script under `uv run --with openpyxl` (no pip install)
  3. --remote → send the grounded spec to Damira's /api/extension/export-xlsx. This
     uploads the spec, so it is opt-in only and never automatic.
  4. none of those → the HTML preview is still written; exit code 4 says the xlsx isn't.

The styling is ported from services/oncall-agent/src/spreadsheet/generator.py so a
workbook looks the same whether the server or the plugin rendered it.
"""

from __future__ import annotations

import argparse
import html
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import workbook_grounding  # noqa: E402 — sibling, stdlib only

EXIT_OK = 0
EXIT_ERROR = 1
EXIT_NO_XLSX = 4
_REEXEC_ENV = "DAMIRA_WORKBOOK_CHILD"
_CWD_ENV = "DAMIRA_WORKBOOK_CWD"
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")

THEME = {
    "header_bg": "1F4E79", "header_font": "FFFFFF", "input_cell": "FFFFCC",
    "auto_filled": "F2F2F2", "alt_row": "F5F9FF", "pass": "C6EFCE", "fail": "FFC7CE",
    "warning": "FFEB9C",
}


class SpecError(Exception):
    pass


# ---------------------------------------------------------------------------
# Spec
# ---------------------------------------------------------------------------

def normalise(data) -> dict:
    """Validate the shape the host model wrote; fill defaults. Raises SpecError."""
    if not isinstance(data, dict):
        raise SpecError("workbook spec must be a JSON object")
    sheets = data.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        raise SpecError("workbook spec needs a non-empty 'sheets' list")
    used = set()
    out = []
    for i, s in enumerate(sheets):
        if not isinstance(s, dict) or not isinstance(s.get("headers"), list) or not s["headers"]:
            raise SpecError(f"sheet {i + 1} needs a non-empty 'headers' list")
        name = _BAD_SHEET_CHARS.sub("-", str(s.get("name") or f"Sheet{i + 1}"))[:31] or f"Sheet{i + 1}"
        while name.lower() in used or name.lower() == "summary":
            name = (name[:28] + f"-{i + 1}")[:31]
        used.add(name.lower())
        out.append({
            "name": name,
            "headers": [str(h) for h in s["headers"]],
            "rows": [r for r in (s.get("rows") or []) if isinstance(r, dict)],
            "input_columns": list(s.get("input_columns") or []),
            "dropdowns": {str(k): list_source(v) for k, v in dict(s.get("dropdowns") or {}).items()},
            "conditional_rules": [r for r in (s.get("conditional_rules") or []) if isinstance(r, dict)],
            "auto_id_prefix": str(s.get("auto_id_prefix") or ""),
            "formulas": dict(s.get("formulas") or {}),
            "url_columns": list(s.get("url_columns") or []),
        })
    return {
        "title": str(data.get("title") or "Workbook"),
        "summary_fields": {str(k): v for k, v in dict(data.get("summary_fields") or {}).items()},
        "executive_summary": str(data.get("executive_summary") or ""),
        "key_metrics": {str(k): v for k, v in dict(data.get("key_metrics") or {}).items()},
        "sheets": out,
    }


def list_source(options) -> str:
    """Dropdown options go inside a quoted Excel list formula ("a,b,c"). A '"' in them
    would end the string and let the rest run as a formula, so quotes and control
    characters are dropped. Excel caps a list source at 255 characters."""
    return re.sub(r'["\x00-\x1f]', "", str(options))[:255]


def cell_value(row: dict, header: str):
    key = header.lower().replace(" ", "_")
    # "Time Est." keys as time_est. or time_est; accept both.
    value = row.get(key, "") or row.get(header, "") or row.get(re.sub(r"\W+", "_", header.lower()).strip("_"), "")
    return "" if value is None else value


def _formula(template: str, headers: "list[str]", row_idx: int) -> str:
    from_letter = {h: _col_letter(i + 1) for i, h in enumerate(headers)}
    for h, letter in from_letter.items():
        template = template.replace("{" + h + "}", f"{letter}{row_idx}")
    return template


# Functions a workbook formula may call. Anything else (WEBSERVICE, HYPERLINK, IMPORT*,
# DDE, external [book] refs) could send cell contents off the machine when the file is
# opened, so such a formula is written as plain text instead.
_SAFE_FUNCS = {
    "IF", "IFS", "IFERROR", "AND", "OR", "NOT", "SUM", "SUMIF", "SUMIFS", "COUNT", "COUNTA",
    "COUNTIF", "COUNTIFS", "COUNTBLANK", "AVERAGE", "AVERAGEIF", "MIN", "MAX", "ROUND",
    "ROUNDUP", "ROUNDDOWN", "ABS", "INT", "MOD", "PRODUCT", "LEN", "TRIM", "UPPER", "LOWER",
    "LEFT", "RIGHT", "MID", "FIND", "SEARCH", "SUBSTITUTE", "EXACT", "CONCAT", "CONCATENATE",
    "TEXTJOIN", "TEXT", "VALUE", "TODAY", "NOW", "DATE", "DATEDIF", "NETWORKDAYS", "ISBLANK",
    "ISNUMBER", "ISERROR", "ISTEXT", "VLOOKUP", "XLOOKUP", "INDEX", "MATCH",
}
_FUNC_CALL = re.compile(r"([A-Za-z_][A-Za-z0-9_.]*)\s*\(")


def safe_formula(formula: str) -> bool:
    if not formula.startswith("=") or any(ch in formula for ch in "[]|!"):
        return False
    body = re.sub(r'"[^"]*"', '""', formula)  # string literals can't call anything
    return all(name.upper() in _SAFE_FUNCS for name in _FUNC_CALL.findall(body))


def _col_letter(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def row_cells(sheet: dict) -> "list[list]":
    """Resolved cell values per data row (auto IDs and formulas applied)."""
    rows = []
    for row_idx, row in enumerate(sheet["rows"], start=2):
        cells = []
        for col_idx, header in enumerate(sheet["headers"], start=1):
            if col_idx == 1 and sheet["auto_id_prefix"]:
                cells.append(f"{sheet['auto_id_prefix']}-{row_idx - 1:03d}")
            elif header in sheet["formulas"]:
                cells.append(_formula(str(sheet["formulas"][header]), sheet["headers"], row_idx))
            else:
                cells.append(cell_value(row, header))
        rows.append(cells)
    return rows


# ---------------------------------------------------------------------------
# xlsx (openpyxl, imported lazily)
# ---------------------------------------------------------------------------

def render_xlsx(spec: dict, out_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
    from openpyxl.worksheet.datavalidation import DataValidation

    def fill(c):
        return PatternFill(start_color=c, end_color=c, fill_type="solid")

    header_fill, input_fill = fill(THEME["header_bg"]), fill(THEME["input_cell"])
    auto_fill, alt_fill = fill(THEME["auto_filled"]), fill(THEME["alt_row"])
    green, red, amber = fill(THEME["pass"]), fill(THEME["fail"]), fill(THEME["warning"])
    red_font = Font(bold=True, color="9C0006")
    side = Side(style="thin", color="B4C6E7")
    border = Border(left=side, right=side, top=side, bottom=side)
    wrap = Alignment(wrap_text=True, vertical="top")

    def text(ws, row, column, value):
        """Cell text from the model is data: a leading '=' must not become a formula."""
        c = ws.cell(row=row, column=column, value=value)
        if isinstance(value, str) and value.startswith("="):
            c.data_type = "s"
        return c

    def protect(ws):
        ws.protection.sheet = True
        ws.protection.sort = False
        ws.protection.autoFilter = False

    def equal_rules(ws, rng, greens, reds, ambers=()):
        for v in greens:
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{v}"'], fill=green))
        for v in reds:
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{v}"'],
                                                          fill=red, font=red_font))
        for v in ambers:
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{v}"'], fill=amber))

    wb = Workbook()
    wb.remove(wb.active)

    # Summary
    ws = wb.create_sheet("Summary")
    ws.merge_cells("A1:D1")
    text(ws, 1, 1, spec["title"])
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color=THEME["header_font"])
    ws["A1"].fill = header_fill
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 40
    fields = {"Generated Date": date.today().isoformat(), **spec["summary_fields"]}
    r = 3
    for label, value in fields.items():
        text(ws, r, 1, label).font = Font(bold=True)
        ws.cell(row=r, column=1).border = border
        c = text(ws, r, 2, value)
        c.border = border
        if isinstance(value, str) and value.startswith("["):
            c.fill, c.protection = input_fill, Protection(locked=False)
        else:
            c.fill = auto_fill
        r += 1
    r += 1
    if spec["key_metrics"]:
        ws.cell(row=r, column=1, value="Key Metrics").font = Font(bold=True, size=12)
        r += 1
        for metric, value in spec["key_metrics"].items():
            text(ws, r, 1, metric).border = border
            text(ws, r, 2, value).border = border
            r += 1
    if spec["executive_summary"]:
        r += 1
        ws.cell(row=r, column=1, value="Executive Summary").font = Font(bold=True, size=12)
        r += 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r + 4, end_column=4)
        c = text(ws, r, 1, spec["executive_summary"])
        c.alignment, c.border = wrap, border
    for col, width in zip("ABCD", (20, 40, 20, 20)):
        ws.column_dimensions[col].width = width
    protect(ws)

    # Data sheets
    for sheet in spec["sheets"]:
        ws = wb.create_sheet(sheet["name"])
        headers = sheet["headers"]
        letters = {h: _col_letter(i + 1) for i, h in enumerate(headers)}
        for i, h in enumerate(headers, start=1):
            c = text(ws, 1, i, h)
            c.fill, c.border = header_fill, border
            c.font = Font(name="Calibri", size=11, bold=True, color=THEME["header_font"])
            c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 25

        for row_idx, cells in enumerate(row_cells(sheet), start=2):
            for col_idx, value in enumerate(cells, start=1):
                header = headers[col_idx - 1]
                is_formula = header in sheet["formulas"] and safe_formula(str(value))
                c = (ws.cell(row=row_idx, column=col_idx, value=value) if is_formula
                     else text(ws, row_idx, col_idx, value))
                c.border, c.alignment = border, wrap
                if header in sheet["input_columns"]:
                    c.fill, c.protection = input_fill, Protection(locked=False)
                else:
                    c.fill = alt_fill if row_idx % 2 == 0 else auto_fill

        last = max(len(sheet["rows"]) + 1, 200)
        for header, options in sheet["dropdowns"].items():
            if header in letters:
                dv = DataValidation(type="list", formula1=f'"{list_source(options)}"', allow_blank=True,
                                    showErrorMessage=True)
                dv.error, dv.errorTitle = "Please select from the dropdown", "Invalid Entry"
                ws.add_data_validation(dv)
                dv.add(f"{letters[header]}2:{letters[header]}{last}")

        rules = list(sheet["conditional_rules"])
        if workbook_grounding.GROUNDING_HEADER in letters:
            rules.append({"column": workbook_grounding.GROUNDING_HEADER, "type": "grounding"})
        for rule in rules:
            letter = letters.get(rule.get("column", ""))
            if not letter:
                continue
            rng = f"{letter}2:{letter}{last}"
            kind = rule.get("type", "")
            if kind == "status":
                equal_rules(ws, rng, ["Complete"], ["Failed", "Rolled Back"], ["In Progress"])
            elif kind == "pass_fail":
                equal_rules(ws, rng, ["Pass"], ["Fail"])
            elif kind == "yes_no":
                equal_rules(ws, rng, ["Yes"], ["No"])
            elif kind == "grounding":
                equal_rules(ws, rng, [workbook_grounding.MATCHED], [workbook_grounding.UNVERIFIED])
            elif kind == "risk_score":
                ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThanOrEqual", formula=["15"],
                                                              fill=red, font=red_font))
                ws.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["6", "14"], fill=amber))
                ws.conditional_formatting.add(rng, CellIsRule(operator="lessThanOrEqual", formula=["5"], fill=green))

        for header in sheet["url_columns"]:
            if header in letters:
                col = headers.index(header) + 1
                for row in range(2, ws.max_row + 1):
                    c = ws.cell(row=row, column=col)
                    if c.value and str(c.value).startswith(("http://", "https://")):
                        c.hyperlink = str(c.value)
                        c.font = Font(color="0563C1", underline="single")

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{_col_letter(len(headers))}1"
        ws.print_title_rows = "1:1"
        for i, h in enumerate(headers, start=1):
            longest = max([len(h)] + [len(str(ws.cell(row=r, column=i).value or ""))
                                      for r in range(2, min(ws.max_row + 1, 50))])
            ws.column_dimensions[_col_letter(i)].width = max(12, min(50, longest + 2))
        protect(ws)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    _replace_into(Path(out_path), lambda tmp: wb.save(tmp))




# ---------------------------------------------------------------------------
# index.html — self-contained preview (stdlib). Base: generator.workbook_to_html.
# ---------------------------------------------------------------------------

_GREEN = {"ok", "pass", "complete", "yes", "resolved", "done", "source-matched"}
_RED = {"fail", "failed", "error", "no", "rolled back", "unverified"}
_AMBER = {"warning", "in progress", "pending", "partial"}

_CSS = """
:root{--bg:#f7f8fa;--panel:#fff;--ink:#1b2430;--muted:#5b6675;--line:#d9dee6;--head:#1f4e79;
--head-ink:#fff;--alt:#f3f6fb;--ok:#1a7f45;--ok-bg:#dff3e6;--bad:#a4161a;--bad-bg:#fbe1e1;
--warn:#8a6100;--warn-bg:#fff1c7;--input:#fffbe0}
@media (prefers-color-scheme:dark){:root{--bg:#0f1318;--panel:#171c23;--ink:#e6eaf0;--muted:#9aa5b4;
--line:#2a323d;--head:#1d3b5c;--alt:#1b212a;--ok:#5fd394;--ok-bg:#15301f;--bad:#ff8a8a;--bad-bg:#3a1719;
--warn:#f0c85a;--warn-bg:#352a0e;--input:#2b2912}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);
font:14px/1.45 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,sans-serif}
main{max-width:1280px;margin:0 auto;padding:20px 16px 48px}h1{font-size:22px;margin:0 0 4px}
.sub{color:var(--muted);margin:0 0 16px}.stats{display:flex;flex-wrap:wrap;gap:10px;margin:0 0 16px}
.stat{background:var(--panel);border:1px solid var(--line);border-radius:8px;padding:10px 14px;min-width:120px}
.stat b{display:block;font-size:20px}.stat span{color:var(--muted);font-size:12px}
.tabs{display:flex;flex-wrap:wrap;gap:6px;margin:0 0 12px}
.tabs button{font:inherit;border:1px solid var(--line);background:var(--panel);color:var(--ink);
border-radius:6px;padding:6px 12px;cursor:pointer}.tabs button[aria-selected=true]{background:var(--head);
color:var(--head-ink);border-color:var(--head)}section{display:none}section.on{display:block}
.wrap{overflow-x:auto;background:var(--panel);border:1px solid var(--line);border-radius:8px}
table{border-collapse:collapse;width:100%}th,td{padding:6px 10px;border-bottom:1px solid var(--line);
text-align:left;vertical-align:top;white-space:pre-wrap}th{background:var(--head);color:var(--head-ink);
position:sticky;top:0;cursor:pointer;white-space:nowrap}tbody tr:nth-child(even){background:var(--alt)}
td.in{background:var(--input)}td.ok{color:var(--ok);background:var(--ok-bg);font-weight:600}
td.bad{color:var(--bad);background:var(--bad-bg);font-weight:600}
td.warn{color:var(--warn);background:var(--warn-bg);font-weight:600}
dl{display:grid;grid-template-columns:max-content 1fr;gap:6px 16px;background:var(--panel);
border:1px solid var(--line);border-radius:8px;padding:14px;margin:0 0 12px}dt{font-weight:600}dd{margin:0}
.exec{background:var(--panel);border:1px solid var(--line);border-radius:8px;padding:14px;white-space:pre-wrap}
code{font-family:ui-monospace,SFMono-Regular,Menlo,monospace;font-size:13px}
"""

_JS = """
document.querySelectorAll('.tabs button').forEach(function(b){b.addEventListener('click',function(){
document.querySelectorAll('.tabs button').forEach(function(x){x.setAttribute('aria-selected','false')});
document.querySelectorAll('section').forEach(function(s){s.classList.remove('on')});
b.setAttribute('aria-selected','true');document.getElementById(b.dataset.t).classList.add('on');});});
document.querySelectorAll('th').forEach(function(th){th.addEventListener('click',function(){
var t=th.closest('table'),i=Array.prototype.indexOf.call(th.parentNode.children,th),
b=t.tBodies[0],rows=Array.prototype.slice.call(b.rows),asc=th.dataset.dir!=='asc';th.dataset.dir=asc?'asc':'desc';
rows.sort(function(a,c){var x=a.cells[i].textContent,y=c.cells[i].textContent,
n=parseFloat(x)-parseFloat(y);return (isNaN(n)?x.localeCompare(y):n)*(asc?1:-1)});
rows.forEach(function(r){b.appendChild(r)});});});
"""


def _status_class(value) -> str:
    v = str(value).strip().lower()
    if v in _GREEN:
        return "ok"
    if v in _RED:
        return "bad"
    if v in _AMBER:
        return "warn"
    return ""


def render_html(spec: dict, grounding: dict) -> str:
    e = html.escape
    rows_total = sum(len(s["rows"]) for s in spec["sheets"])
    stats = [(str(len(spec["sheets"])), "sheets"), (str(rows_total), "rows")]
    if grounding.get("checked"):
        stats.append((f"{grounding['pct']}%", "commands matched evidence"))
        stats.append((str(grounding["unverified"]), "UNVERIFIED"))
    for sheet in spec["sheets"]:
        if "Status" in sheet["headers"] and sheet["rows"]:
            done = sum(1 for r in sheet["rows"] if str(cell_value(r, "Status")).lower() in ("complete", "done"))
            stats.append((f"{done}/{len(sheet['rows'])}", f"{sheet['name']} complete"))

    out = ["<!doctype html><html lang=\"en\"><head><meta charset=\"utf-8\">",
           "<meta name=\"viewport\" content=\"width=device-width,initial-scale=1\">",
           f"<title>{e(spec['title'])}</title><style>{_CSS}</style></head><body><main>",
           f"<h1>{e(spec['title'])}</h1>",
           f"<p class=\"sub\">Generated {date.today().isoformat()} by the Damira plugin. "
           "The .xlsx next to this file is the working copy.</p>",
           "<div class=\"stats\">"]
    out += [f"<div class=\"stat\"><b>{e(v)}</b><span>{e(label)}</span></div>" for v, label in stats]
    out.append("</div><div class=\"tabs\" role=\"tablist\">")
    tabs = [("s0", "Summary")] + [(f"s{i + 1}", s["name"]) for i, s in enumerate(spec["sheets"])]
    for i, (tid, name) in enumerate(tabs):
        out.append(f"<button role=\"tab\" data-t=\"{tid}\" aria-selected=\"{'true' if i == 0 else 'false'}\">"
                   f"{e(name)}</button>")
    out.append("</div>")

    out.append("<section id=\"s0\" class=\"on\"><dl>")
    for label, value in {"Generated": date.today().isoformat(), **spec["summary_fields"],
                         **spec["key_metrics"]}.items():
        out.append(f"<dt>{e(str(label))}</dt><dd>{e(str(value))}</dd>")
    out.append("</dl>")
    if spec["executive_summary"]:
        out.append(f"<div class=\"exec\">{e(spec['executive_summary'])}</div>")
    out.append("</section>")

    for i, sheet in enumerate(spec["sheets"]):
        out.append(f"<section id=\"s{i + 1}\"><div class=\"wrap\"><table><thead><tr>")
        out += [f"<th>{e(h)}</th>" for h in sheet["headers"]]
        out.append("</tr></thead><tbody>")
        for cells in row_cells(sheet):
            out.append("<tr>")
            for header, value in zip(sheet["headers"], cells):
                classes = []
                if header in sheet["input_columns"]:
                    classes.append("in")
                sc = _status_class(value)
                if sc:
                    classes.append(sc)
                attr = f" class=\"{' '.join(classes)}\"" if classes else ""
                text = e(str(value))
                if header.strip().lower() in workbook_grounding.GROUNDING_HEADER_KEYS and value:
                    text = f"<code>{text}</code>"
                out.append(f"<td{attr}>{text}</td>")
            out.append("</tr>")
        out.append("</tbody></table></div></section>")

    out.append(f"</main><script>{_JS}</script></body></html>")
    return "".join(out)


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------

def _openpyxl_available() -> bool:
    return importlib.util.find_spec("openpyxl") is not None


def _replace_into(target: Path, write) -> None:
    """Write via a temp file in the same folder, then rename over the target. A symlink
    planted at the target is replaced, never followed, so no other file gets overwritten."""
    fd, tmp = tempfile.mkstemp(prefix=".damira-", suffix=target.suffix, dir=str(target.parent))
    os.close(fd)
    try:
        write(tmp)
        umask = os.umask(0)
        os.umask(umask)
        os.chmod(tmp, 0o666 & ~umask)  # mkstemp makes 0600; a normal file follows the umask
        os.replace(tmp, str(target))
    except BaseException:
        try:
            os.unlink(tmp)
        except OSError:
            pass
        raise


def _write_bytes(target: Path, data: bytes) -> None:
    def write(tmp: str) -> None:
        with open(tmp, "wb") as fh:
            fh.write(data)
    _replace_into(target, write)


def _outputs(spec_path: Path, out_dir: "Path | None") -> "tuple[Path, Path]":
    folder = out_dir or spec_path.parent
    stem = spec_path.name
    for suffix in (".damira.json", ".json"):
        if stem.endswith(suffix):
            stem = stem[: -len(suffix)]
            break
    return folder / f"{stem or 'workbook'}.xlsx", folder / "index.html"


def render(spec_path, out_dir=None, evidence_dir=None, xlsx=True, ground=True, host="cli",
           redact=False) -> dict:
    """Ground, write index.html, and (if xlsx) the workbook in-process. Needs openpyxl for xlsx.

    redact=True (layer 3, before an artifact publish) replaces hosts, IPs, secrets and
    people with tokens after grounding, and never writes an xlsx."""
    started = time.monotonic()
    spec_path = Path(spec_path)
    try:
        raw = json.loads(spec_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise SpecError(f"cannot read {spec_path}: {exc}")
    spec = normalise(raw)

    evidence = Path(evidence_dir) if evidence_dir else spec_path.parent / "evidence"
    corpus = workbook_grounding.load_corpus(evidence) if ground else ""
    unverified = checked = 0
    if ground:
        spec, unverified, checked = workbook_grounding.ground_spec(spec, corpus)
    grounding = {"checked": checked, "unverified": unverified,
                 "pct": workbook_grounding.grounded_pct(unverified, checked),
                 "line": workbook_grounding.summary_line(unverified, checked, not corpus) if ground else ""}

    redactions = {}
    if redact:
        import redact_spec

        spec, redactions = redact_spec.redact(spec)
        xlsx = False

    xlsx_path, html_path = _outputs(spec_path, Path(out_dir) if out_dir else None)
    html_path.parent.mkdir(parents=True, exist_ok=True)
    _write_bytes(html_path, render_html(spec, grounding).encode("utf-8"))
    written = None
    if xlsx:
        render_xlsx(spec, xlsx_path)
        written = xlsx_path

    _log(host, "redacted" if redact else ("ok" if written or not xlsx else "partial"), spec, grounding, started)
    return {"xlsx": written, "xlsx_target": xlsx_path, "html": html_path, "grounding": grounding,
            "spec": spec, "redactions": redactions}


def _log(host: str, status: str, spec: dict, grounding: dict, started: float) -> None:
    try:
        import events
        events.emit({
            "event": "deliverable_generated", "deliverable_type": "workbook", "host": host,
            "status": status, "grounded_pct": grounding.get("pct"),
            "duration_ms": int((time.monotonic() - started) * 1000),
            "counts": {"sheets": len(spec["sheets"]), "rows": sum(len(s["rows"]) for s in spec["sheets"]),
                       "checked": grounding.get("checked", 0), "unverified": grounding.get("unverified", 0)},
        })
    except Exception:  # noqa: BLE001 — logging never breaks rendering
        pass


def _reexec_with_uv(argv: "list[str]") -> "int | None":
    uv = shutil.which("uv")
    if not uv or os.environ.get(_REEXEC_ENV):
        return None
    print("openpyxl is not installed; running the renderer with uv (fetches openpyxl 3.1.5 once).",
          file=sys.stderr)
    # --no-config and a pinned interpreter, run from an empty temp folder: a workspace
    # .python-version, uv.toml or [tool.uv] must not pick the interpreter or the index.
    # The child changes back to the caller's folder so relative paths still work.
    cmd = [uv, "run", "--quiet", "--no-project", "--no-config", "--python", sys.executable,
           "--with", "openpyxl==3.1.5", "python", str(Path(__file__).resolve()), *argv]
    env = {**os.environ, _REEXEC_ENV: "1", _CWD_ENV: os.getcwd()}
    try:
        with tempfile.TemporaryDirectory(prefix="damira-uv-") as neutral:
            return subprocess.run(cmd, cwd=neutral, env=env, timeout=600).returncode
    except (OSError, subprocess.SubprocessError):
        return None


def remote_payload(spec: dict) -> dict:
    """What --remote uploads. The server writes a leading '=' as a live formula, so the
    same rule as the local renderer is applied before the spec leaves: text starting
    with '=' gets a leading space, formulas outside the allowlist are dropped, and
    dropdown options lose their quotes."""
    def text(node):
        if isinstance(node, str):
            return " " + node if node.startswith("=") else node
        if isinstance(node, list):
            return [text(n) for n in node]
        if isinstance(node, dict):
            return {text(k) if isinstance(k, str) else k: text(v) for k, v in node.items()}
        return node

    out = text({k: v for k, v in spec.items() if k != "sheets"})
    out["sheets"] = []
    for sheet in spec.get("sheets") or []:
        formulas = {k: v for k, v in (sheet.get("formulas") or {}).items()
                    if isinstance(v, str) and safe_formula(v)}
        clean = text({k: v for k, v in sheet.items() if k not in ("formulas", "dropdowns")})
        clean["formulas"] = formulas
        clean["dropdowns"] = {k: list_source(v) for k, v in (sheet.get("dropdowns") or {}).items()}
        out["sheets"].append(clean)
    return out


def _remote_xlsx(spec: dict, target: Path) -> None:
    """Opt-in fallback: the gateway renders the (already grounded) spec with the same engine."""
    import urllib.error
    import urllib.request

    import damira

    key, is_demo = damira.resolve_key()
    if is_demo:
        raise SpecError("--remote needs your own Damira API key (the shared demo key can't export). "
                        "Install uv instead to render locally: https://docs.astral.sh/uv/")
    api = os.environ.get("DAMIRA_API_URL", damira.DEFAULT_API_URL).rstrip("/")
    req = urllib.request.Request(
        f"{api}/api/extension/export-xlsx", data=json.dumps(remote_payload(spec)).encode("utf-8"),
        method="POST",
        headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json",
                 "User-Agent": damira.USER_AGENT, "Accept": _XLSX_MIME})
    try:
        with urllib.request.urlopen(req, timeout=120) as resp:
            body = resp.read()
    except urllib.error.HTTPError as exc:
        raise SpecError(f"remote export failed: HTTP {exc.code}")
    except urllib.error.URLError as exc:
        raise SpecError(f"remote export failed: {exc.reason}")
    if not body.startswith(b"PK"):
        raise SpecError("remote export did not return an .xlsx file")
    target.parent.mkdir(parents=True, exist_ok=True)
    _write_bytes(target, body)


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(prog="damira workbook",
                                description="Render workbook.json to .xlsx + index.html, with a grounding check")
    p.add_argument("spec", help="WorkbookSpec JSON, e.g. documents/<id>/workbook.json")
    p.add_argument("--out-dir", default="", help="default: the spec's folder")
    p.add_argument("--evidence", default="", help="evidence folder (default: <spec dir>/evidence)")
    p.add_argument("--html-only", action="store_true", help="write only the index.html preview")
    p.add_argument("--no-grounding", action="store_true", help=argparse.SUPPRESS)
    p.add_argument("--redact", action="store_true",
                   help="write a redacted preview only (hosts, IPs, secrets, people → tokens); "
                        "use with --out-dir before publishing it anywhere")
    p.add_argument("--remote", action="store_true",
                   help="no local openpyxl/uv: send the spec to Damira to render the .xlsx (uploads it)")
    p.add_argument("--host", default=os.environ.get("DAMIRA_HOST", "cli"),
                   choices=["cli", "claude-code", "cursor"], help=argparse.SUPPRESS)
    return p


def main(argv: "list[str] | None" = None) -> int:
    argv = list(sys.argv[1:] if argv is None else argv)
    if os.environ.get(_REEXEC_ENV) and os.environ.get(_CWD_ENV):
        try:
            os.chdir(os.environ[_CWD_ENV])  # re-exec'd from a neutral folder; back to the caller's
        except OSError:
            pass
    a = build_parser().parse_args(argv)

    if a.redact:
        a.html_only, a.remote = True, False
    local_xlsx = not a.html_only and not a.remote and _openpyxl_available()
    if not a.html_only and not a.remote and not local_xlsx:
        code = _reexec_with_uv(argv)
        if code is not None:
            return code

    try:
        result = render(a.spec, a.out_dir or None, a.evidence or None, xlsx=local_xlsx,
                        ground=not a.no_grounding, host=a.host, redact=a.redact)
        if a.remote and not a.html_only:
            _remote_xlsx(result["spec"], result["xlsx_target"])
            result["xlsx"] = result["xlsx_target"]
    except SpecError as exc:
        print(f"damira workbook: {exc}", file=sys.stderr)
        return EXIT_ERROR
    except OSError as exc:
        print(f"damira workbook: cannot write the output: {exc.strerror or exc}", file=sys.stderr)
        return EXIT_ERROR

    if result["xlsx"]:
        print(f"Workbook: {result['xlsx']}")
    print(f"Preview:  {result['html']}")
    if result["grounding"]["line"]:
        print(result["grounding"]["line"])
    if a.redact:
        done = ", ".join(f"{n} {k}" for k, n in sorted(result["redactions"].items())) or "nothing matched"
        print(f"Redacted: {done}. Check the preview for anything the patterns missed before sharing it.")
    if not result["xlsx"] and not a.html_only:
        print("No .xlsx written: this machine has neither openpyxl nor uv. Install uv "
              "(https://docs.astral.sh/uv/) and run this again, or ask the user before re-running "
              "with --remote, which sends the spec to Damira to render.", file=sys.stderr)
        return EXIT_NO_XLSX
    return EXIT_OK


if __name__ == "__main__":
    sys.exit(main())
