#!/usr/bin/env python3
# /// script
# requires-python = ">=3.9"
# dependencies = ["python-docx==1.2.0"]
# ///
"""Convert a MOP, change control or incident report (markdown) to .docx + .html (#485).

The host model writes the markdown; the .md stays the source of truth. This writes the
Word copy for CAB submission and a self-contained preview next to it:

    damira docx documents/<id>/mop.md            # → mop.docx + mop.html
    uv run md_to_docx.py documents/<id>/mop.md   # directly (PEP 723)

The .docx needs python-docx, and the plugin stays stdlib-only, so python-docx is
imported lazily and found in this order (the render_workbook.py pattern, #484):
  1. already importable in this interpreter → convert in-process
  2. `uv` on PATH → re-run this script under `uv run --with python-docx` (no pip install)
  3. --remote → send the markdown to Damira's /api/extension/export-docx. This uploads
     the document, so it is opt-in only and never automatic.
  4. none of those → the HTML preview is still written; exit code 4 says the .docx isn't.

The Word styling is ported from services/oncall-agent/src/docx_converter.py so a
document looks the same whether the server or the plugin converted it.
"""

from __future__ import annotations

import argparse
import html
import importlib.util
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

EXIT_OK = 0
EXIT_ERROR = 1
EXIT_NO_DOCX = 4
_PIN = "python-docx==1.2.0"
_REEXEC_ENV = "DAMIRA_DOCX_CHILD"
_CWD_ENV = "DAMIRA_DOCX_CWD"
_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MAX_BYTES = 1_000_000  # a MOP is text; anything larger is not one (the gateway caps the same)

_H = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_BULLET = re.compile(r"^\s*[-*+]\s+(.+)")
_NUMBERED = re.compile(r"^\s*\d+[.)]\s+(.+)")
_RULE = re.compile(r"^\s*([-*_])(\s*\1){2,}\s*$")
_TABLE_SEP = re.compile(r"^\s*\|?[\s:|-]+\|[\s:|-]*$")
_INLINE = re.compile(r"`([^`]+)`|\*\*([^*]+)\*\*|__([^_]+)__")


class DocError(Exception):
    pass


def _plain(text: str) -> str:
    """Inline markup stripped, for headings and the document title."""
    return _INLINE.sub(lambda m: m.group(1) or m.group(2) or m.group(3), text).strip()


def _cells(line: str) -> "list[str]":
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def parse(md: str) -> "list[tuple]":
    """Markdown → blocks: ("h", level, text) ("p", text) ("ul"|"ol", [items])
    ("table", [rows]) ("code", text) ("hr",). Shared by the .docx and HTML writers so
    both show the same document."""
    blocks: "list[tuple]" = []
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith("```") or s.startswith("~~~"):
            fence, body = s[:3], []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith(fence):
                body.append(lines[i])
                i += 1
            blocks.append(("code", "\n".join(body)))
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not _TABLE_SEP.match(lines[i]):
                    rows.append(_cells(lines[i]))
                i += 1
            if rows:
                blocks.append(("table", rows))
            continue
        h = _H.match(s)
        if h:
            blocks.append(("h", len(h.group(1)), h.group(2)))
        elif _RULE.match(s):
            blocks.append(("hr",))
        elif _BULLET.match(line) or _NUMBERED.match(line):
            kind = "ul" if _BULLET.match(line) else "ol"
            pat = _BULLET if kind == "ul" else _NUMBERED
            items = []
            while i < len(lines) and pat.match(lines[i]):
                items.append(pat.match(lines[i]).group(1).strip())
                i += 1
            blocks.append((kind, items))
            continue
        elif s:
            para = [s]
            while (i + 1 < len(lines) and lines[i + 1].strip() and not lines[i + 1].strip().startswith(("|", "```", "~~~", "#"))
                   and not _BULLET.match(lines[i + 1]) and not _NUMBERED.match(lines[i + 1]) and not _RULE.match(lines[i + 1])):
                i += 1
                para.append(lines[i].strip())
            blocks.append(("p", " ".join(para)))
        i += 1
    return blocks


def doc_title(md: str, fallback: str) -> str:
    for block in parse(md):
        if block[0] == "h" and block[1] == 1:
            return _plain(block[2])
    return fallback.replace("-", " ").replace("_", " ").strip().title() or "Document"


# ---------------------------------------------------------------------------
# HTML preview (#484 layer 2): stdlib, escaped, no external assets, no script
# ---------------------------------------------------------------------------

_CSS = """
:root{--bg:#f7f8fa;--panel:#fff;--ink:#1b2430;--muted:#5b6675;--line:#d9dee6;--head:#1f4e79;
--head-ink:#fff;--alt:#f3f6fb;--code:#f2f2f2}
@media (prefers-color-scheme:dark){:root{--bg:#0f1318;--panel:#171c23;--ink:#e6eaf0;--muted:#9aa5b4;
--line:#2a323d;--head:#1d3b5c;--alt:#1b212a;--code:#10151b}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);
font:15px/1.55 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,sans-serif}
main{max-width:900px;margin:0 auto;padding:24px 16px 56px;background:var(--panel);border-left:1px solid var(--line);
border-right:1px solid var(--line);min-height:100vh}h1{font-size:24px;margin:0 0 4px}h2{font-size:19px;
margin:28px 0 8px;padding-bottom:4px;border-bottom:1px solid var(--line)}h3{font-size:16px;margin:20px 0 6px}
.sub{color:var(--muted);margin:0 0 20px;font-size:13px}.wrap{overflow-x:auto;margin:10px 0}
table{border-collapse:collapse;width:100%}th,td{border:1px solid var(--line);padding:6px 10px;text-align:left;
vertical-align:top}th{background:var(--head);color:var(--head-ink)}tbody tr:nth-child(even){background:var(--alt)}
pre{background:var(--code);border:1px solid var(--line);border-radius:6px;padding:10px 12px;overflow-x:auto}
code{font-family:ui-monospace,SFMono-Regular,Menlo,Consolas,monospace;font-size:13px}
hr{border:0;border-top:1px solid var(--line);margin:20px 0}
"""


def _inline_html(text: str) -> str:
    out, pos = [], 0
    for m in _INLINE.finditer(text):
        out.append(html.escape(text[pos:m.start()]))
        if m.group(1) is not None:
            out.append(f"<code>{html.escape(m.group(1))}</code>")
        else:
            out.append(f"<strong>{html.escape(m.group(2) or m.group(3))}</strong>")
        pos = m.end()
    out.append(html.escape(text[pos:]))
    return "".join(out)


def render_html(md: str, title: str) -> str:
    e = html.escape
    out = ["<!doctype html><html lang=\"en\"><head><meta charset=\"utf-8\">",
           "<meta name=\"viewport\" content=\"width=device-width,initial-scale=1\">",
           f"<title>{e(title)}</title><style>{_CSS}</style></head><body><main>",
           f"<p class=\"sub\">Generated {date.today().isoformat()} by the Damira plugin. "
           "The .docx next to this file is the copy to submit; the .md is the source.</p>"]
    for block in parse(md):
        kind = block[0]
        if kind == "h":
            out.append(f"<h{block[1]}>{_inline_html(block[2])}</h{block[1]}>")
        elif kind == "p":
            out.append(f"<p>{_inline_html(block[1])}</p>")
        elif kind in ("ul", "ol"):
            out.append(f"<{kind}>" + "".join(f"<li>{_inline_html(x)}</li>" for x in block[1]) + f"</{kind}>")
        elif kind == "code":
            out.append(f"<pre><code>{e(block[1])}</code></pre>")
        elif kind == "hr":
            out.append("<hr>")
        elif kind == "table":
            rows = block[1]
            width = max(len(r) for r in rows)
            out.append("<div class=\"wrap\"><table><thead><tr>")
            out += [f"<th>{_inline_html(c)}</th>" for c in rows[0] + [""] * (width - len(rows[0]))]
            out.append("</tr></thead><tbody>")
            for r in rows[1:]:
                out.append("<tr>" + "".join(f"<td>{_inline_html(c)}</td>" for c in r + [""] * (width - len(r))) + "</tr>")
            out.append("</tbody></table></div>")
    out.append("</main></body></html>")
    return "".join(out)


# ---------------------------------------------------------------------------
# .docx (needs python-docx; styling ported from the agent's docx_converter.py)
# ---------------------------------------------------------------------------

def markdown_to_docx(md: str, out_path: Path, title: str) -> None:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    def shade(props, fill: str) -> None:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        props.append(shd)

    def borders(cell) -> None:
        # Inside <w:tcBorders>, before <w:shd>: Word rejects edges placed straight in tcPr.
        tc_pr = cell._tc.get_or_add_tcPr()
        box = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            tag = OxmlElement(f"w:{edge}")
            tag.set(qn("w:val"), "single")
            tag.set(qn("w:sz"), "4")
            tag.set(qn("w:color"), "000000")
            box.append(tag)
        tc_pr.append(box)

    def runs(para, text: str, bold: bool = False) -> None:
        pos = 0
        for m in list(_INLINE.finditer(text)) + [None]:
            plain = text[pos:m.start()] if m else text[pos:]
            if plain:
                r = para.add_run(plain)
                r.bold = bold or None
                r.font.name = "Calibri"
            if m is None:
                break
            if m.group(1) is not None:
                r = para.add_run(m.group(1))
                r.bold = bold or None
                r.font.name = "Consolas"
                r.font.size = Pt(10)
            else:
                r = para.add_run(m.group(2) or m.group(3))
                r.bold = True
                r.font.name = "Calibri"
            pos = m.end()

    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(1)
        section.left_margin = section.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    doc.core_properties.title = title
    doc.core_properties.author = "Damira"

    sizes = {1: 16, 2: 14, 3: 12}
    for block in parse(md):
        kind = block[0]
        if kind == "h":
            para = doc.add_heading(_plain(block[2]), level=min(block[1], 9))
            for r in para.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(sizes.get(block[1], 11))
        elif kind == "p":
            runs(doc.add_paragraph(), block[1])
        elif kind in ("ul", "ol"):
            for item in block[1]:
                runs(doc.add_paragraph(style="List Bullet" if kind == "ul" else "List Number"), item)
        elif kind == "hr":
            doc.add_paragraph()
        elif kind == "code":
            for line in block[1].splitlines() or [""]:
                para = doc.add_paragraph()
                shade(para._p.get_or_add_pPr(), "F2F2F2")
                para.paragraph_format.left_indent = Inches(0.25)
                para.paragraph_format.space_after = Pt(0)
                r = para.add_run(line or " ")
                r.font.name = "Consolas"
                r.font.size = Pt(10)
        elif kind == "table":
            rows = block[1]
            width = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=width)
            table.style = "Table Grid"
            for ri, cells in enumerate(rows):
                for ci in range(width):
                    cell = table.cell(ri, ci)
                    para = cell.paragraphs[0]
                    runs(para, cells[ci] if ci < len(cells) else "", bold=ri == 0)
                    borders(cell)
                    if ri == 0:
                        shade(cell._tc.get_or_add_tcPr(), "D9E1F2")
    _replace_into(out_path, lambda tmp: doc.save(tmp))


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------

def _docx_available() -> bool:
    return importlib.util.find_spec("docx") is not None


def _replace_into(target: Path, write) -> None:
    """Write via a temp file in the same folder, then rename over the target. A symlink
    planted at the target is replaced, never followed."""
    fd, tmp = tempfile.mkstemp(prefix=".damira-", suffix=target.suffix, dir=str(target.parent))
    os.close(fd)
    try:
        write(tmp)
        umask = os.umask(0)
        os.umask(umask)
        os.chmod(tmp, 0o666 & ~umask)
        os.replace(tmp, str(target))
    except BaseException:
        try:
            os.unlink(tmp)
        except OSError:
            pass
        raise


def _write_bytes(target: Path, data: bytes) -> None:
    def write(tmp: str) -> None:
        with open(tmp, "wb") as fh:
            fh.write(data)
    _replace_into(target, write)


def convert(md_path, out_dir=None, docx=True, redact=False, title=None, host="cli") -> dict:
    """Write <stem>.html, and <stem>.docx in-process when docx=True (needs python-docx).

    redact=True (layer 3, before an artifact publish) replaces hosts, IPs and secrets with
    tokens and never writes the .docx."""
    started = time.monotonic()
    md_path = Path(md_path)
    try:
        if md_path.stat().st_size > MAX_BYTES:
            raise DocError(f"{md_path} is over {MAX_BYTES // 1_000_000} MB; that is not a MOP")
        md = md_path.read_text(encoding="utf-8")
    except (OSError, UnicodeDecodeError) as exc:
        raise DocError(f"cannot read {md_path}: {exc}")
    if not md.strip():
        raise DocError(f"{md_path} is empty")

    redactions = {}
    if redact:
        import redact_spec

        md, redactions = redact_spec.redact_markdown(md)
        docx = False
    title = title or doc_title(md, md_path.stem)

    folder = Path(out_dir) if out_dir else md_path.parent
    folder.mkdir(parents=True, exist_ok=True)
    stem = md_path.stem or "document"
    html_path, docx_path = folder / f"{stem}.html", folder / f"{stem}.docx"
    _write_bytes(html_path, render_html(md, title).encode("utf-8"))
    written = None
    if docx:
        markdown_to_docx(md, docx_path, title)
        written = docx_path

    _log(host, "redacted" if redact else ("ok" if written or not docx else "partial"), md, started)
    return {"docx": written, "docx_target": docx_path, "html": html_path, "markdown": md,
            "title": title, "redactions": redactions}


def _log(host: str, status: str, md: str, started: float) -> None:
    try:
        import events
        blocks = parse(md)
        events.emit({
            "event": "deliverable_generated", "deliverable_type": "docx", "host": host, "status": status,
            "duration_ms": int((time.monotonic() - started) * 1000),
            "counts": {k: sum(1 for b in blocks if b[0] == k) for k in ("h", "table", "code")},
        })
    except Exception:  # noqa: BLE001 — logging never breaks conversion
        pass


def _reexec_with_uv(argv: "list[str]") -> "int | None":
    uv = shutil.which("uv")
    if not uv or os.environ.get(_REEXEC_ENV):
        return None
    print(f"python-docx is not installed; converting with uv (fetches {_PIN} once).", file=sys.stderr)
    # Pinned interpreter, --no-config, run from an empty folder: a workspace
    # .python-version or uv.toml must not pick the interpreter or the index.
    cmd = [uv, "run", "--quiet", "--no-project", "--no-config", "--python", sys.executable,
           "--with", _PIN, "python", str(Path(__file__).resolve()), *argv]
    env = {**os.environ, _REEXEC_ENV: "1", _CWD_ENV: os.getcwd()}
    try:
        with tempfile.TemporaryDirectory(prefix="damira-uv-") as neutral:
            return subprocess.run(cmd, cwd=neutral, env=env, timeout=600).returncode
    except (OSError, subprocess.SubprocessError):
        return None


def _remote_docx(md: str, title: str, target: Path) -> None:
    """Opt-in fallback: the gateway converts with the agent's docx_converter."""
    import urllib.error
    import urllib.request

    import damira

    key, is_demo = damira.resolve_key()
    if is_demo:
        raise DocError("--remote needs your own Damira API key (the shared demo key can't export). "
                       "Install uv instead to convert locally: https://docs.astral.sh/uv/")
    api = os.environ.get("DAMIRA_API_URL", damira.DEFAULT_API_URL).rstrip("/")
    req = urllib.request.Request(
        f"{api}/api/extension/export-docx",
        data=json.dumps({"markdown": md, "title": title}).encode("utf-8"), method="POST",
        headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json",
                 "User-Agent": damira.USER_AGENT, "Accept": _DOCX_MIME})
    try:
        with urllib.request.urlopen(req, timeout=120) as resp:
            body = resp.read()
    except urllib.error.HTTPError as exc:
        raise DocError(f"remote export failed: HTTP {exc.code}")
    except urllib.error.URLError as exc:
        raise DocError(f"remote export failed: {exc.reason}")
    if not body.startswith(b"PK"):
        raise DocError("remote export did not return a .docx file")
    _write_bytes(target, body)


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(prog="damira docx",
                                description="Convert a MOP / change control / incident report (.md) "
                                            "to .docx + a self-contained .html preview")
    p.add_argument("markdown", help="e.g. documents/<id>/mop.md")
    p.add_argument("--title", default="", help="document title (default: the first # heading)")
    p.add_argument("--out-dir", default="", help="default: the markdown file's folder")
    p.add_argument("--html-only", action="store_true", help="write only the .html preview")
    p.add_argument("--redact", action="store_true",
                   help="write a redacted preview only (hosts, IPs, secrets → tokens); "
                        "use with --out-dir before publishing it anywhere")
    p.add_argument("--remote", action="store_true",
                   help="no local python-docx/uv: send the markdown to Damira to convert (uploads it)")
    p.add_argument("--host", default=os.environ.get("DAMIRA_HOST", "cli"),
                   choices=["cli", "claude-code", "cursor"], help=argparse.SUPPRESS)
    return p


def main(argv: "list[str] | None" = None) -> int:
    argv = list(sys.argv[1:] if argv is None else argv)
    if os.environ.get(_REEXEC_ENV) and os.environ.get(_CWD_ENV):
        try:
            os.chdir(os.environ[_CWD_ENV])
        except OSError:
            pass
    a = build_parser().parse_args(argv)

    if a.redact:
        a.html_only, a.remote = True, False
    local = not a.html_only and not a.remote and _docx_available()
    if not a.html_only and not a.remote and not local:
        code = _reexec_with_uv(argv)
        if code is not None:
            return code

    try:
        result = convert(a.markdown, a.out_dir or None, docx=local, redact=a.redact,
                         title=a.title or None, host=a.host)
        if a.remote and not a.html_only:
            _remote_docx(result["markdown"], result["title"], result["docx_target"])
            result["docx"] = result["docx_target"]
    except DocError as exc:
        print(f"damira docx: {exc}", file=sys.stderr)
        return EXIT_ERROR
    except OSError as exc:
        print(f"damira docx: cannot write the output: {exc.strerror or exc}", file=sys.stderr)
        return EXIT_ERROR

    if result["docx"]:
        print(f"Word:     {result['docx']}")
    print(f"Preview:  {result['html']}")
    if a.redact:
        done = ", ".join(f"{n} {k}" for k, n in sorted(result["redactions"].items())) or "nothing matched"
        print(f"Redacted: {done}. Check the preview for anything the patterns missed before sharing it.")
    if not result["docx"] and not a.html_only:
        print("No .docx written: this machine has neither python-docx nor uv. Install uv "
              "(https://docs.astral.sh/uv/) and run this again, or ask the user before re-running "
              "with --remote, which sends the document to Damira to convert.", file=sys.stderr)
        return EXIT_NO_DOCX
    return EXIT_OK


if __name__ == "__main__":
    sys.exit(main())
